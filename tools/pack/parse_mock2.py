import docx,json,re
U='src/'
F={4:'APPSC_Geography_MockTest4_ExamPattern_Bilingual.docx',5:'APPSC_History_MockTest1_ExamPattern_Bilingual.docx'}
TE=re.compile('[ఀ-౿]')
def cells(row):
    seen=[];out=[]
    for c in row.cells:
        if c._tc in seen: continue
        seen.append(c._tc); out.append(c)
    return out
def paras(c): return [p.text.strip() for p in c.paragraphs if p.text.strip()]
out={}
for n,f in F.items():
    d=docx.Document(U+f); qs=[]; key={}
    for t in d.tables:
        rows=t.rows
        c0=cells(rows[0])
        p0=paras(c0[0]) if c0 else []
        if p0 and re.match(r'^\d+\.\s',p0[0]):
            m=re.match(r'^(\d+)\.\s*(.*)',p0[0]); num=int(m.group(1)); p0[0]=m.group(2)
            ti=next(i for i,x in enumerate(p0) if TE.search(x))
            en,te=p0[:ti],p0[ti:]
            opts=[]
            for r in rows[1:]:
                for c in cells(r):
                    ps=paras(c)
                    if not ps: continue
                    e=re.sub(r'^\(\d\)\s*','',ps[0]); tt=re.sub(r'^\(\d\)\s*','',ps[1]) if len(ps)>1 else e
                    opts.append([e,tt])
            assert len(opts)==4,(n,num,opts)
            if abs(len(en)-len(te))>0: print('lines differ',n,num,len(en),len(te))
            qs.append({'n':num,'en':en,'te':te,'o':opts})
        else:
            for r in rows:
                cs=[c.text.strip() for c in cells(r)]
                if len(cs)>=4 and re.match(r'^\d+$',cs[0]) and re.match(r'^[1-4]$',cs[1]): key[int(cs[0])]=cs[1:4]
    for q in qs:
        k=key[q['n']]; q['a']=int(k[0])-1; q['xe']=k[1]; q['xt']=k[2]
    assert [q['n'] for q in qs]==list(range(1,51)),n
    out[n]={'q':qs}; print(n,len(qs),len(key))
    from collections import Counter; print(' answer spread',Counter(q['a']+1 for q in qs))
json.dump(out,open('newmock.json','w',encoding='utf-8'),ensure_ascii=False)
